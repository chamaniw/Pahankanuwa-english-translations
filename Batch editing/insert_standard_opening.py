from docx import Document
from pathlib import Path
import pandas as pd

FOLDER = Path(r"C:\Users\chama\OneDrive\Pahankanuwa Project\Editing English Translations")

NEW_PARA_3 = "(Pahankaṇuwa Dhamma Sermon No. )"

NEW_PARA_4 = """Namo tassa bhagavato arahato sammāsambuddhassa

(Homage to the Blessed One, the Worthy One, the Fully Self-Enlightened One)"""

changes = []

for file in FOLDER.rglob("*.docx"):

    try:

        doc = Document(file)

        paras = [p.text for p in doc.paragraphs]

        if len(paras) < 2:
            print(f"Skipped: {file.name}")
            continue

        # Prevent accidental double insertion
        if NEW_PARA_3 in paras or "Namo tassa bhagavato arahato sammāsambuddhassa" in "\n".join(paras):
            print(f"Already updated: {file.name}")
            continue

        new_para_list = []

        for i, text in enumerate(paras):

            new_para_list.append(text)

            # After paragraph 2 (index 1)
            if i == 1:
                new_para_list.append(NEW_PARA_3)
                new_para_list.append(NEW_PARA_4)

        # rebuild document
        new_doc = Document()

        for p in new_para_list:
            new_doc.add_paragraph(p)

        new_doc.save(file)

        changes.append({
            "File": file.name,
            "Inserted": "Yes"
        })

        print(f"Updated: {file.name}")

    except Exception as e:
        print(f"Error: {file.name}")
        print(e)

report = pd.DataFrame(changes)

report.to_excel(
    FOLDER / "Paragraph_Insertion_Report.xlsx",
    index=False
)

print()
print("DONE")
print(f"Files updated: {len(changes)}")